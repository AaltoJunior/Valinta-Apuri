import json
import os
import shutil

import filetype
import piexif
import piexif.helper
from openpyxl import load_workbook
from openpyxl_image_loader import SheetImageLoader
from PIL import Image


def load_images_from_excel(file_path):
    tmp_folder = "static/img_tmp"
    cur_folder = "static/img_cur"

    if os.path.exists(tmp_folder):
        shutil.rmtree(tmp_folder)
    os.makedirs(tmp_folder, exist_ok=True)
    os.makedirs(cur_folder, exist_ok=True)

    workbook = load_workbook(file_path)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError(f"'{file_path}' has no sheet named 'Sheet1' (found: {workbook.sheetnames})")
    sheet = workbook["Sheet1"]
    image_loader = SheetImageLoader(sheet)
    # A broken image only loses that one card's picture; web.py serves
    # generic.jpg for missing images.
    for row in range(2, sheet.max_row + 1):
        cell = f"H{row}"
        image_name = row - 2
        try:
            if image_loader.image_in(cell):
                image = image_loader.get(cell)
                image.save(os.path.join(tmp_folder, f"{image_name}.png"))
                print(f"Saved image in {cell} as {tmp_folder}/{image_name}.png")
        except Exception as error:
            print(f"Skipping image in {cell}: {error}")

    for image_file in os.listdir(tmp_folder):
        image_path = os.path.join(tmp_folder, image_file)
        try:
            if os.path.isfile(image_path) and filetype.is_image(image_path):
                _convert_image(image_path)
        except Exception as error:
            print(f"Skipping image {image_file}, conversion failed: {error}")
            if os.path.exists(image_path):
                os.remove(image_path)

    for entry in os.scandir(cur_folder):
        if entry.is_dir(follow_symlinks=False):
            shutil.rmtree(entry.path)
        else:
            os.unlink(entry.path)
    shutil.copytree(tmp_folder, cur_folder, dirs_exist_ok=True)
    print("Pictures reloaded successfully!")


def _convert_image(image_path):
    with Image.open(image_path) as image:
        webp_path = image_path.rsplit(".", 1)[0] + ".webp"
        width, height = image.size
        max_side = 500
        if height >= width:
            new_height = max_side
            new_width = max(1, int(width * (new_height / height)))
        else:
            new_width = max_side
            new_height = max(1, int(height * (new_width / width)))

        exif = {"0th": {}, "Exif": {}, "1st": {}, "GPS": {}}
        exif["Exif"][piexif.ExifIFD.UserComment] = piexif.helper.UserComment.dump(
            json.dumps({"ow": width, "oh": height, "of": image.format})
        )
        image.resize((new_width, new_height), Image.Resampling.LANCZOS).convert("RGB").save(
            webp_path,
            format="WEBP",
            exif=piexif.dump(exif),
            lossless=False,
            quality=90,
            method=6,
        )
    os.remove(image_path)
    print(
        f"Converted {os.path.basename(image_path)} to {webp_path} "
        f"og w,h = {width}x{height} -> {new_width}x{new_height}"
    )
